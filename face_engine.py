import face_recognition
import numpy as np
import json
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

# Store data OUTSIDE OneDrive to prevent sync from deleting/renaming files mid-write.
DATA_DIR = os.path.abspath(os.path.join(os.path.expanduser("~"), "face_attendance_data"))
USERS_FILE = os.path.join(DATA_DIR, "users.json")
ATTENDANCE_FILE = os.path.join(DATA_DIR, "attendance.xlsx")


def _ensure_data_dir():
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
    except (FileExistsError, OSError):
        pass


_ensure_data_dir()


def load_users():
    if not os.path.exists(USERS_FILE):
        return {}
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            content = f.read().strip()
            if not content:
                return {}
            return json.loads(content)
    except (json.JSONDecodeError, UnicodeDecodeError, OSError):
        backup = USERS_FILE + ".corrupt"
        try:
            os.replace(USERS_FILE, backup)
        except OSError:
            pass
        return {}


def save_users(users):
    _ensure_data_dir()
    tmp_file = USERS_FILE + ".tmp"
    with open(tmp_file, "w", encoding="utf-8") as f:
        json.dump(users, f, indent=2, ensure_ascii=False)
    os.replace(tmp_file, USERS_FILE)


def encode_face_from_image(image_path):
    """Extract 128-D encoding from image. Returns list or None."""
    img = face_recognition.load_image_file(image_path)
    locations = face_recognition.face_locations(img, model="hog")
    if not locations:
        return None
    encodings = face_recognition.face_encodings(img, locations)
    if not encodings:
        return None
    return encodings[0].tolist()


def match_face(unknown_encoding, users, tolerance=0.5):
    if not users:
        return None, 0.0

    all_encodings = []
    all_names = []
    for name, info in users.items():
        for enc in info["encodings"]:
            all_encodings.append(np.array(enc))
            all_names.append(name)

    if not all_encodings:
        return None, 0.0

    distances = face_recognition.face_distance(
        all_encodings, np.array(unknown_encoding)
    )
    best_idx = int(np.argmin(distances))
    best_distance = float(distances[best_idx])

    if best_distance <= tolerance:
        confidence = round((1 - best_distance) * 100, 1)
        return all_names[best_idx], confidence
    return None, 0.0


def log_attendance(name, student_id, confidence):
    _ensure_data_dir()
    today = datetime.now().strftime("%Y-%m-%d")
    time_now = datetime.now().strftime("%H:%M:%S")

    if os.path.exists(ATTENDANCE_FILE):
        wb = load_workbook(ATTENDANCE_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Date", "Time", "Name", "Student ID", "Confidence"])

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == today and row[3] == student_id:
            return False

    ws.append([today, time_now, name, student_id, f"{confidence}%"])
    wb.save(ATTENDANCE_FILE)
    return True


def get_all_attendance():
    _ensure_data_dir()
    if not os.path.exists(ATTENDANCE_FILE):
        return []
    wb = load_workbook(ATTENDANCE_FILE)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({
            "date": row[0],
            "time": row[1],
            "name": row[2],
            "student_id": row[3],
            "confidence": row[4],
        })
    return rows